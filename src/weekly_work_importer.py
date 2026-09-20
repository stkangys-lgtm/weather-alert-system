"""건설사업회의 XLSX에서 현장명과 금주 진행공종만 추출한다.

원본의 담당자, 연락처, 공사금액, 수금액, 현안은 읽거나 Google Sheets에 올리지 않는다.
"""

from __future__ import annotations

import argparse
import re
from datetime import datetime

from openpyxl import load_workbook

from src import settings as config


WEEKLY_WORKSHEET = "주간공종"
WEEKLY_WORK_HEADER = ["현장명", "진행공종"]

# 회의자료의 표기와 관제시스템 표기가 달라도 추측 매칭하지 않도록 승인된 별칭만 관리한다.
SITE_ALIASES = {
    "오리온 진천신공장": ["오리온 진천 신공장"],
    "오리온 진천 기숙사": ["오리온 진천 기숙사"],
    "연희·연남동 공공주택": ["연희연남 공공주택"],
    "청정고원 스포츠센터": ["청정고원 스포츠센터"],
    "LX 논현 업무시설": ["LX 논현사옥"],
    "군포복합개발": ["군포 복합개발"],
    "오리온수협 목포 김공장": ["오리온수협 목포 김공장"],
    "렉서스 동탄 네트워크": ["렉서스 동탄 네트워크", "렉서스 동탄"],
    "화천군부대 시설공사": ["화천 군부대 숙소", "화천 군부대"],
    "반얀트리호텔 근생동": ["반얀트리 증축공사", "반얀트리호텔 근생동"],
    "양산 부산대병원": ["양산 부산대병원", "양산부산대병원"],
    "포항~안동2 국도건설공사": ["포항~안동2 (1공구) 국도"],
    "시흥능곡 주변도로": ["시흥능곡 주변도로 관리"],
    "뇌죽천 하천재해예방": ["뇌죽천 하천재해예방"],
    "산솔면 하수처리장": ["산솔면 하수처리장", "산솔면 하수처리"],
    "단월정수장 시설공사": ["단월정수장 현대화사업"],
    "후포 공공하수처리": ["울진군 후포 공공하수증설"],
    "동해안 바닷가 자동차길": ["강릉 바닷가 자동차길 조성"],
    "풍각지구 정비사업": ["풍각지구 종합정비사업"],
    "영주 가흥정수장": ["영주시 가흥정수장"],
    "송산그린시티 용수공급시설": ["송산그린시티 용수공급시설"],
}


def _normalize(value):
    return re.sub(r"[^0-9A-Za-z가-힣]", "", str(value or "")).lower()


def _clean_work(value):
    lines = []
    for line in str(value or "").replace("\r", "").split("\n"):
        cleaned = re.sub(r"[ \t]+", " ", line).strip()
        if cleaned:
            lines.append(cleaned)
    return "\n".join(lines)


def _parse_period(header):
    text = str(header or "")
    dates = re.findall(r"(\d{2})[./-](\d{1,2})[./-](\d{1,2})", text)
    if len(dates) < 2:
        return None
    start = datetime(2000 + int(dates[0][0]), int(dates[0][1]), int(dates[0][2])).date()
    end = datetime(2000 + int(dates[1][0]), int(dates[1][1]), int(dates[1][2])).date()
    return start, end


def _find_current_block(ws):
    candidates = []
    for row in range(1, min(ws.max_row, 15) + 1):
        site_cols = []
        for col in range(1, min(ws.max_column, 60) + 1):
            value = re.sub(r"\s+", "", str(ws.cell(row, col).value or ""))
            if value == "현장명":
                site_cols.append(col)
            if "주요공정-금주" not in value:
                continue
            left_sites = [site_col for site_col in site_cols if site_col < col]
            if not left_sites:
                continue
            period = _parse_period(ws.cell(row, col).value)
            candidates.append((period[0] if period else datetime.min.date(), row, max(left_sites), col, period))
    if not candidates:
        raise ValueError(f"{ws.title}: '현장명/주요공정-금주' 헤더를 찾지 못했습니다")
    _, header_row, site_col, work_col, period = max(candidates, key=lambda item: item[0])
    return header_row, site_col, work_col, period


def _extract_source_records(workbook):
    records = []
    periods = []
    target_sheets = [ws for ws in workbook.worksheets if "주요현장" in ws.title]
    if not target_sheets:
        raise ValueError("건축/토목 주요현장 시트를 찾지 못했습니다")

    for ws in target_sheets:
        header_row, site_col, work_col, period = _find_current_block(ws)
        if period:
            periods.append(period)
        order_col = site_col - 1
        for row in range(header_row + 1, ws.max_row + 1):
            order = ws.cell(row, order_col).value
            site_name = ws.cell(row, site_col).value
            work = _clean_work(ws.cell(row, work_col).value)
            if not isinstance(order, (int, float)) or not site_name or not work:
                continue
            records.append({
                "source_site_name": str(site_name).strip(),
                "work": work,
                "source_sheet": ws.title,
                "source_row": row,
                "source_hidden": bool(ws.row_dimensions[row].hidden),
            })
    period = max(periods, key=lambda item: item[0]) if periods else None
    return records, period


def _alias_index(active_sites):
    index = {}
    for site in active_sites:
        canonical = site["site_name"]
        aliases = [canonical] + SITE_ALIASES.get(canonical, [])
        for alias in aliases:
            key = _normalize(alias)
            if key in index and index[key] != canonical:
                raise ValueError(f"중복 현장 별칭: {alias}")
            index[key] = canonical
    return index


def parse_weekly_work(path, active_sites=None):
    """원본 XLSX를 읽어 승인된 활성 현장만 매칭한 결과와 검증정보를 반환한다."""
    active_sites = active_sites if active_sites is not None else config.SITES
    workbook = load_workbook(path, data_only=True, read_only=False)
    source_records, period = _extract_source_records(workbook)
    aliases = _alias_index(active_sites)

    matched = {}
    unmatched_source = []
    for record in source_records:
        canonical = aliases.get(_normalize(record["source_site_name"]))
        if not canonical:
            unmatched_source.append(record["source_site_name"])
            continue
        if canonical in matched:
            raise ValueError(f"같은 현장이 회의자료에 중복되었습니다: {canonical}")
        matched[canonical] = record

    active_names = [site["site_name"] for site in active_sites]
    rows = [[name, matched[name]["work"]] for name in active_names if name in matched]
    missing_active = [name for name in active_names if name not in matched]
    return {
        "rows": rows,
        "period": period,
        "matched_count": len(rows),
        "missing_active_sites": missing_active,
        "unmatched_source_sites": unmatched_source,
    }


def upload_weekly_work(path):
    # XLSX 분석만 할 때는 Google Sheets 패키지와 인증이 필요하지 않도록 업로드 시점에 불러온다.
    from src.sheets_client import get_worksheet, replace_rows

    report = parse_weekly_work(path)
    worksheet = get_worksheet(
        config.GOOGLE_SHEETS_SPREADSHEET_ID,
        WEEKLY_WORKSHEET,
        credentials_path=config.GOOGLE_SHEETS_CREDENTIALS_PATH,
        credentials_json=config.GOOGLE_SERVICE_ACCOUNT_JSON,
    )
    replace_rows(worksheet, WEEKLY_WORK_HEADER, report["rows"])
    return report


def main():
    parser = argparse.ArgumentParser(description="건설사업회의 XLSX의 금주 공종을 추출합니다")
    parser.add_argument("xlsx_path")
    parser.add_argument("--upload", action="store_true", help="Google Sheets 주간공종 탭을 교체합니다")
    args = parser.parse_args()
    report = upload_weekly_work(args.xlsx_path) if args.upload else parse_weekly_work(args.xlsx_path)
    print(f"[주간공종] 매칭 {report['matched_count']}개 현장")
    if report["period"]:
        print(f"[주간공종] 적용기간 {report['period'][0]} ~ {report['period'][1]}")
    if report["missing_active_sites"]:
        print(f"[주간공종] 회의자료 누락: {', '.join(report['missing_active_sites'])}")
    if report["unmatched_source_sites"]:
        print(f"[주간공종] 관제 제외: {', '.join(report['unmatched_source_sites'])}")


if __name__ == "__main__":
    main()
